from datetime import datetime
import re
from sprnldata.core.frame import ImfFrame
import sprnldata.core.dummy.config as dum
import sprnldata.core.ecos.config as c
from pandaspro.core.frame import FramePro
from openpyxl.utils import column_index_from_string, get_column_letter
import pandas as pd
import imf_datatools
import xlwings as xw
import os
import shutil

# Step 1: read excel of Data Pulling Tool Main Dashboard
# Step 2: retrieve the meta info into a dictionary
# Step 3: parse dictionary and use it with imf_datatools to download data

def _ecosuse(file: str = None, debug: bool = False):
    wb = xw.Book(file)
    ws = wb.sheets['Dashboard']

    sectioncol = ['B', 'F', 'J', 'N', 'R']
    df = pd.DataFrame()
    df['COUNTRY'] = ''
    df['dates'] = ''
    errmsg = ''

    for index, col in zip(range(5), sectioncol):
        dbname = ws.range(f'{col}3').value
        colindex = column_index_from_string(col)
        if dbname is not None:
            freq = ws.range(f'{col}4').value
            start = ws.range(f'{col}5').value
            end = ws.range(f'{col}6').value
            countryselect = ws.range(f'{col}7').value

            if countryselect == 'All countries w/o aggregates':
                clist = dum.wo_aggregate
            elif countryselect == 'All countries w aggregates (WLD, WAEMU, EURO, and ECCU)':
                clist = dum.w_aggregate
            else:
                clist_colindex = colindex - 1
                clist_col = get_column_letter(clist_colindex)
                clist_raw = [int(c) for c in ws.range(f'{clist_col}10:{clist_col}209').value if
                             c is not None and str(c).strip() != '']
                clist_ifs = [c for c in clist_raw if
                             isinstance(c, (int, float)) or (isinstance(c, str) and c.isdigit())]
                clist_iso = [c for c in clist_raw if isinstance(c, str) and re.match('[A-Z]{3}', c)]
                clist_name = [c for c in clist_raw if c not in (clist_ifs, clist_iso)]

            counterlist = [c for c in ws.range(f'{col}10:{col}209').value if c is not None]
            indlist_colindex = colindex + 1
            indlist_col = get_column_letter(indlist_colindex)
            indlist = [i for i in ws.range(f'{indlist_col}10:{indlist_col}209').value if i is not None]

            if debug:
                print(
                    f'Section {index + 1}/5: imf_datatools.get_ecos_sdmx_data({dbname}, {clist}, {indlist}, freq={freq}, longformat=True)')
            data = imf_datatools.get_ecos_sdmx_data(dbname, clist, indlist, freq=freq, longformat=True)

            ## print notification message including time duration for pulling the data

            if data is not None:
                if start is not None:
                    data = data[data['dates'].dt.year >= start]
                if end is not None:
                    data = data[data['dates'].dt.year <= end]
                df = pd.merge(df, data, on=['COUNTRY', 'dates'], how='outer')

            else:
                errmsg = errmsg + '\n' + f'{dbname} {indlist} NO data available' + '\n'
        else:
            print(f'Section {index + 1}/5: Skipped as Database (row3) not provided')
            continue

    print(errmsg)
    df.rename(columns={'COUNTRY': 'ifscode'}, inplace=True)
    df['ifscode'] = df['ifscode'].astype(int)
    df['year'] = df['dates'].dt.year
    df.drop(columns='dates', inplace=True)
    new_order = ['ifscode', 'year'] + [col for col in df.columns if
                                       col not in ['ifscode', 'year']]  # reorder the columns
    return df[new_order]


class EcosData(ImfFrame):
    def __init__(self, data=None, debug: bool = False, *args, **kwargs):
        if isinstance(data, (pd.DataFrame, FramePro)):
            super().__init__(data=data, *args, **kwargs)
        else:
            result = _ecosuse(file=data, debug=debug)
            super().__init__(data=result, *args, **kwargs)


class EcosSet:
    def __init__(self, workbook: str = 'Data Pulling Template.xlsx'):
        destination_file = os.path.join(os.getcwd(), workbook)

        if os.path.exists(destination_file):
            print(f'Opening current {workbook} in the folder ... ')
        else:
            shutil.copyfile(c.template, destination_file)
            print(f'{workbook} copied from template into the folder ... ')
        self.path = os.path.join(os.getcwd(), workbook)
        xw.Book(self.path)

    def pull(self, debug: bool = False):
        op = EcosData(
            data=self.path,
            debug=debug
        )
        return op
