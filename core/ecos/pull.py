from swxl.pandaspro.core import pwread
from openpyxl.utils import column_index_from_string, get_column_letter
import pandas as pd
import imf_datatools
import xlwings as xw
import core.dummy.config as c

template = r'C:\ProgramData\anaconda3\Lib\sprnldata\core\ecos\dev-template.xlsx'
wb = xw.Book(template)
ws = wb.sheets['Dashboard']

sectioncol = ['B', 'F', 'J', 'N', 'R']
df = pd.DataFrame()
df['COUNTRY'] = ''
df['dates']=''
errmsg = ''

for col in sectioncol:
    dbname = ws.range(f'{col}3').value
    colindex = column_index_from_string(col)
    if dbname is not None:
        freq = ws.range(f'{col}4').value
        start = ws.range(f'{col}5').value
        end = ws.range(f'{col}6').value
        countryselect = ws.range(f'{col}7').value

        if countryselect == 'All countries w/o aggregates':
            clist = c.wo_aggregate
        elif countryselect == 'All countries w aggregates (WLD, WAEMU, EURO, and ECCU)':
            clist = c.w_aggregate
        else:
            clist_colindex = colindex - 1
            clist_col = get_column_letter(clist_colindex)
            clist = pwread(template, 'Dashboard', f'{clist_col}10:{clist_col}209', firstrow=False)[0][clist_col.lower()].to_list()

        counterlist = pwread(template, 'Dashboard', f'{col}10:{col}209', firstrow=False)[0][col.lower()].to_list()
        indllist_colindex =  colindex + 1
        indllist_col = get_column_letter(indllist_colindex)
        indlist = pwread(template, 'Dashboard', f'{indllist_col}10:{indllist_col}209', firstrow=False)[0][indllist_col.lower()].dropna().to_list()

        data = imf_datatools.get_ecos_sdmx_data(dbname, clist, indlist, freq=freq, longformat=True)
        if data is not None:
            if start is not None:
                data = data[data['dates'].dt.year >= start]
            if end is not None:
                data = data[data['dates'].dt.year <= end]
            df = pd.merge(df, data, on=['COUNTRY', 'dates'], how='outer')
        else:
            errmsg = errmsg + dbname + indlist + 'not available'
    else:
        continue

df.rename(columns={'COUNTRY':'ifscode'}, inplace=True)
df['ifscode'] = df['ifscode'].astype(int)
final = pd.merge(df, c.dummy, on='ifscode', how='left')