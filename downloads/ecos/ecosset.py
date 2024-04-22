import sprnldata.core.dummy.config as dum
import sprnldata.downloads.ecos.config as c
from sprnldata.utils.myecosuse import myecosuse
from openpyxl.utils import column_index_from_string, get_column_letter
import xlwings as xw
import os
import shutil
import datetime
import time
from sprnldata.downloads import ecos_root

folder_ecos = f'{ecos_root}/Ecos'

# Step 1: read excel of Data Pulling Tool Main Dashboard
# Step 2: retrieve the meta info into a dictionary
# Step 3: parse dictionary and use it with imf_datatools to download data

def _update_log(timestamp, logdict):
    log_entry = f"{timestamp}\n===========================\n{str(logdict)}\n"
    log_file_path = os.path.join(folder_ecos, 'templates', 'log.txt')

    if os.path.exists(log_file_path):
        with open(log_file_path, 'r') as file:
            lines = file.readlines()

        entry_index = None
        end_index = None
        for i, line in enumerate(lines):
            if line.strip() == timestamp:
                entry_index = i
            elif entry_index is not None and line.strip() and line.strip().isdigit() and len(line.strip()) == 8:
                end_index = i - 1
                break

        if entry_index is not None and end_index is not None:
            lines = lines[:entry_index] + [log_entry] + lines[end_index:]
        elif entry_index is not None:
            lines = lines[:entry_index] + [log_entry]
        else:
            lines.append(log_entry)
    else:
        lines = [log_entry]

    with open(log_file_path, 'w') as file:
        file.writelines(lines)


class EcosSet:
    def __init__(
            self,
            workbook: str = 'Data Pulling Template.xlsx'
    ):
        destination_file = os.path.join(os.getcwd(), workbook)
        if os.path.exists(destination_file):
            print(f'Opening current {workbook} in the folder ... ')
        else:
            # noinspection PyUnboundLocalVariable
            shutil.copyfile(c.template, destination_file)
            print(f'{workbook} copied from template into the folder ... ')
        self.path = os.path.join(os.getcwd(), workbook)

        wb = xw.Book(self.path)
        ws = wb.sheets['Dashboard']
        sectioncol = ['B', 'F', 'J', 'N', 'R']

        pull_dict = {}
        for index, col in zip(range(5), sectioncol):
            pull_dict[f'Database {index + 1}'] = {}
            dbname = ws.range(f'{col}3').value
            colindex = column_index_from_string(col)

            freq = ws.range(f'{col}4').value
            start = ws.range(f'{col}5').value
            end = ws.range(f'{col}6').value
            countryselect = ws.range(f'{col}7').value

            if countryselect == 'All countries w/o aggregates':
                clist = dum.wo_aggregate
            elif countryselect == 'All countries w aggregates (WLD, WAEMU, EURO, and ECCU)':
                clist = dum.w_aggregate
            else:
                clist_colindex = colindex - 1
                clist_col = get_column_letter(clist_colindex)
                clist_raw = [int(country) for country in ws.range(f'{clist_col}10:{clist_col}209').value if
                             country is not None and str(country).strip() != '']
                clist_ifs = [country for country in clist_raw if
                             isinstance(country, (int, float)) or (isinstance(country, str) and country.isdigit())]
                #clist_iso = [c for c in clist_raw if isinstance(c, str) and re.match('[A-Z]{3}', c)]
                clist = clist_ifs
                ###############################################
                # Future update: match iso and name to ifscode
                ###############################################

            #counterlist = [c for c in ws.range(f'{col}10:{col}209').value if c is not None]
            indlist_colindex = colindex + 1
            indlist_col = get_column_letter(indlist_colindex)
            indlist = [i for i in ws.range(f'{indlist_col}10:{indlist_col}209').value if i is not None]

            pull_dict[f'Database {index + 1}']['dbname'] = dbname
            pull_dict[f'Database {index + 1}']['clist'] = clist
            pull_dict[f'Database {index + 1}']['indlist'] = indlist
            pull_dict[f'Database {index + 1}']['freq'] = freq
            pull_dict[f'Database {index + 1}']['start'] = start
            pull_dict[f'Database {index + 1}']['end'] = end

        self.pull_dict = pull_dict

    def pull(self, debug: bool = False):
        starttime = time.time()
        data = myecosuse(
            meta_dict=self.pull_dict,
            debug=debug
        )
        data.columns = data.columns.str.lower().str.replace('.a', '', regex=False)

        ### Generate commonly used variables
        data['pfb_gdp'] = (data['ggr'] - data['ggx'] + data['ggei']) / data['ngdp'] * 100  # Primary fiscal balance
        data['iar_bmgs'] = data['iar_bp6'] / (data['bmgs_bp6'] / 12)  # Reserves in months of imports

        timestamp = datetime.datetime.now().strftime('%Y%m%d')
        _update_log(timestamp=timestamp, logdict=self.pull_dict)
        data.to_csv(os.path.join(folder_ecos, f'ecosdata_{timestamp}.csv'), index=False)
        endtime = time.time()
        print(f'Download completed. Time Duration: {round((endtime-starttime)/60, 1)}min')


if __name__ == '__main__':
    a = EcosSet(f'{folder_ecos}/templates/template_20240421.xlsx')
    # a.pull()
